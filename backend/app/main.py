from contextlib import asynccontextmanager
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import re
import secrets
from threading import Event, Thread
from typing import Literal
from uuid import uuid4

from google.auth.exceptions import GoogleAuthError
from docx import Document
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Response, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import case, func, select, text, update
from sqlalchemy.orm import Session, selectinload

from .auth import (
    APP_ENV,
    SESSION_COOKIE,
    SESSION_TTL_SECONDS,
    authenticate_user,
    create_session_token,
    get_current_user,
    hash_password,
    normalize_email,
    seed_demo_user,
    verify_google_credential,
)
from .database import Base, SessionLocal, UPLOAD_DIR, engine, get_db, migrate_screening_schema
from .evaluation import EvaluationCriterion, extract_pdf_text, suggest_criteria, test_ai_service
from .matching import JobRequirement, default_matching_weights
from .models import Resume, ResumeEvaluation, Screening, ScreeningRubric, User
from .worker import run_worker

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_JOB_DESCRIPTION_CHARS = 10_000
MIN_JOB_DESCRIPTION_WORDS = 20
RESUME_EXTENSIONS = {".pdf", ".docx"}
JOB_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_RESUMES_PER_BATCH = 10
MAX_EVALUATION_CRITERIA = 20


class BulkDeleteResumes(BaseModel):
    resume_ids: list[str]


class LoginRequest(BaseModel):
    email: str = Field(min_length=3, max_length=320)
    password: str = Field(min_length=1, max_length=256)


class GoogleSignInRequest(BaseModel):
    credential: str = Field(min_length=1, max_length=8192)


class UpdateScreeningRequest(BaseModel):
    name: str = Field(min_length=2, max_length=160)
    job_description_text: str = Field(min_length=1, max_length=MAX_JOB_DESCRIPTION_CHARS)
    criteria: list[dict] | None = Field(default=None, max_length=MAX_EVALUATION_CRITERIA)


class UpdateRequirementsRequest(BaseModel):
    requirements: list[JobRequirement] = Field(min_length=1, max_length=40)


class ReviewOverrideRequest(BaseModel):
    score: int = Field(ge=0, le=100)
    note: str = Field(min_length=3, max_length=2000)


@asynccontextmanager
async def lifespan(_: FastAPI):
    Base.metadata.create_all(bind=engine)
    migrate_screening_schema(engine)
    with SessionLocal() as db:
        seed_demo_user(db)
    worker_stop = Event()
    worker_thread = None
    if os.getenv("RUN_EVALUATION_WORKER", "true").strip().lower() not in {"0", "false", "no"}:
        worker_thread = Thread(
            target=run_worker,
            kwargs={"stop_event": worker_stop, "initialize": False, "recover": True},
            name="resume-evaluation-worker",
            daemon=True,
        )
        worker_thread.start()
    try:
        yield
    finally:
        worker_stop.set()
        if worker_thread is not None:
            worker_thread.join(timeout=5)


app = FastAPI(
    title="RecruitAI Resume Screening API",
    version="0.3.0",
    description="API for storing resume screening sessions and uploaded documents.",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        origin.strip()
        for origin in os.getenv(
            "CORS_ORIGINS", "http://127.0.0.1:5173,http://localhost:5173"
        ).split(",")
        if origin.strip()
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def count_words(value: str) -> int:
    return len(re.findall(r"\b[\w]+(?:['-][\w]+)*\b", value, flags=re.UNICODE))


def safe_filename(upload: UploadFile) -> str:
    raw_name = (upload.filename or "").replace("\\", "/")
    filename = Path(raw_name).name.strip()
    if not filename or filename in {".", ".."}:
        raise HTTPException(status_code=400, detail="Every uploaded file must have a filename.")
    return filename[:255]


async def read_upload(upload: UploadFile, allowed_extensions: set[str]) -> tuple[str, bytes]:
    filename = safe_filename(upload)
    if Path(filename).suffix.lower() not in allowed_extensions:
        allowed = ", ".join(sorted(allowed_extensions))
        raise HTTPException(status_code=415, detail=f"{filename}: supported file types are {allowed}.")
    content = await upload.read(MAX_FILE_SIZE + 1)
    if not content:
        raise HTTPException(status_code=400, detail=f"{filename}: the file is empty.")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail=f"{filename}: files must be 10 MB or smaller.")
    return filename, content


def extract_job_description(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            return content.decode("utf-8-sig")
        if extension == ".pdf":
            return extract_pdf_text(content)
        if extension == ".docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            return "\n".join(parts)
    except Exception as error:
        raise HTTPException(
            status_code=422,
            detail="Could not read text from the job description file. Use a searchable PDF, DOCX, or plain text file.",
        ) from error
    raise HTTPException(status_code=415, detail="Use a PDF, DOCX, or TXT job description.")


def iso_utc(value: datetime) -> str:
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat()


@app.post("/api/ai/test", tags=["ai"])
def ai_test(_: User = Depends(get_current_user)):
    try:
        model, latency_ms = test_ai_service()
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error
    return {"ok": True, "model": model, "latency_ms": latency_ms}


def screening_payload(screening: Screening) -> dict:
    return {
        "id": screening.id,
        "name": screening.name,
        "created_at": iso_utc(screening.created_at),
        "job_description_text": screening.job_description_text,
        "job_description_filename": screening.job_description_filename,
        "job_requirements": screening.job_requirements,
        "scoring_config": screening.scoring_config,
        "criteria": screening.rubric.criteria if screening.rubric else [],
        "resumes": [
            {
                "id": resume.id,
                "filename": resume.filename,
                "size_bytes": resume.size_bytes,
                "content_type": resume.content_type,
                "created_at": iso_utc(resume.created_at),
                "language_code": resume.language_code,
                "evaluation": (
                    {
                        "status": resume.evaluation.status,
                        "score": resume.evaluation.reviewer_override_score if resume.evaluation.reviewer_override_score is not None else resume.evaluation.score,
                        "model_score": resume.evaluation.score,
                        "job_match_score": resume.evaluation.job_match_score,
                        "preferred_score": resume.evaluation.preferred_score,
                        "bonus_points": resume.evaluation.bonus_points,
                        "required_checks_passed": resume.evaluation.required_checks_passed,
                        "required_checks": resume.evaluation.required_checks,
                        "job_evidence": resume.evaluation.job_evidence,
                        "explanation": resume.evaluation.explanation,
                        "assessments": resume.evaluation.assessments,
                        "matching_details": resume.evaluation.matching_details,
                        "reviewer_override_score": resume.evaluation.reviewer_override_score,
                        "reviewer_override_note": resume.evaluation.reviewer_override_note,
                        "reviewer_id": resume.evaluation.reviewer_id,
                        "reviewer_updated_at": iso_utc(resume.evaluation.reviewer_updated_at) if resume.evaluation.reviewer_updated_at else None,
                        "error_message": resume.evaluation.error_message,
                        "attempts": resume.evaluation.attempts,
                        "updated_at": iso_utc(resume.evaluation.updated_at),
                    }
                    if resume.evaluation
                    else None
                ),
            }
            for resume in screening.resumes
            ],
            }


def parse_evaluation_criteria(value: str) -> list[dict]:
    try:
        raw_criteria = json.loads(value)
        if not isinstance(raw_criteria, list):
            raise ValueError("Evaluation criteria must be a list.")
        if len(raw_criteria) > MAX_EVALUATION_CRITERIA:
            raise ValueError(f"Use no more than {MAX_EVALUATION_CRITERIA} evaluation criteria.")
        criteria = [EvaluationCriterion.model_validate(item) for item in raw_criteria]
    except (json.JSONDecodeError, ValidationError, TypeError) as error:
        raise HTTPException(status_code=422, detail="Evaluation criteria are invalid.") from error
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error

    if any(criterion.required for criterion in criteria):
        raise HTTPException(
            status_code=422,
            detail="Additional criteria must be optional preferred qualifications.",
        )

    ids = [criterion.id for criterion in criteria]
    if len(ids) != len(set(ids)):
        raise HTTPException(status_code=422, detail="Evaluation criteria must have unique IDs.")
    return [criterion.model_dump() for criterion in criteria]


async def store_resume_uploads(
    screening: Screening,
    resumes: list[UploadFile],
    uploaded_paths: list[Path],
) -> None:
    if len(resumes) > MAX_RESUMES_PER_BATCH:
        raise HTTPException(
            status_code=413,
            detail=f"Upload no more than {MAX_RESUMES_PER_BATCH} resumes per batch.",
        )
    existing = {
        (resume.filename.casefold(), resume.size_bytes)
        for resume in screening.resumes
    }
    seen_in_batch: set[tuple[str, int]] = set()
    for upload in resumes:
        filename, content = await read_upload(upload, RESUME_EXTENSIONS)
        duplicate_key = (filename.casefold(), len(content))
        if duplicate_key in seen_in_batch:
            raise HTTPException(status_code=422, detail=f"{filename} was included more than once.")
        seen_in_batch.add(duplicate_key)
        if duplicate_key in existing:
            continue

        storage_name = f"{uuid4().hex}{Path(filename).suffix.lower()}"
        file_path = UPLOAD_DIR / storage_name
        file_path.write_bytes(content)
        uploaded_paths.append(file_path)
        screening.resumes.append(
            Resume(
                id=str(uuid4()),
                filename=filename,
                content_type=upload.content_type,
                size_bytes=len(content),
                storage_name=storage_name,
                evaluation=ResumeEvaluation(status="queued"),
            )
        )
        existing.add(duplicate_key)


@app.get("/health", tags=["system"])
def health_check(db: Session = Depends(get_db)) -> dict[str, str]:
    db.execute(text("SELECT 1"))
    return {"status": "ok", "database": "ok"}


@app.post("/api/auth/login", tags=["authentication"])
def login(payload: LoginRequest, response: Response, db: Session = Depends(get_db)) -> dict[str, dict[str, str]]:
    user = authenticate_user(db, payload.email, payload.password)
    if user is None:
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    response.set_cookie(
        key=SESSION_COOKIE,
        value=create_session_token(user),
        max_age=SESSION_TTL_SECONDS,
        httponly=True,
        secure=APP_ENV == "production",
        samesite="strict",
        path="/",
    )
    return {"user": {"id": user.id, "email": user.email, "full_name": user.full_name}}


@app.post("/api/auth/google", tags=["authentication"])
def google_login(
    payload: GoogleSignInRequest,
    response: Response,
    db: Session = Depends(get_db),
) -> dict[str, dict[str, str]]:
    try:
        claims = verify_google_credential(payload.credential)
    except (GoogleAuthError, ValueError) as error:
        raise HTTPException(status_code=401, detail="Google sign-in could not be verified.") from error

    email = claims.get("email")
    if claims.get("email_verified") is not True or not isinstance(email, str) or "@" not in email:
        raise HTTPException(status_code=401, detail="Google sign-in requires a verified email address.")

    email = normalize_email(email)
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        full_name = claims.get("name")
        user = User(
            id=str(uuid4()),
            email=email,
            full_name=(full_name.strip() if isinstance(full_name, str) and full_name.strip() else email.split("@", 1)[0])[:120],
            password_hash=hash_password(secrets.token_urlsafe(48)),
        )
        db.add(user)
        db.commit()
    elif not user.is_active:
        raise HTTPException(status_code=403, detail="This account is disabled.")

    response.set_cookie(
        key=SESSION_COOKIE,
        value=create_session_token(user),
        max_age=SESSION_TTL_SECONDS,
        httponly=True,
        secure=APP_ENV == "production",
        samesite="strict",
        path="/",
    )
    return {"user": {"id": user.id, "email": user.email, "full_name": user.full_name}}


@app.get("/api/auth/me", tags=["authentication"])
def current_user(user: User = Depends(get_current_user)) -> dict[str, dict[str, str]]:
    return {"user": {"id": user.id, "email": user.email, "full_name": user.full_name}}


@app.post("/api/auth/logout", status_code=status.HTTP_204_NO_CONTENT, tags=["authentication"])
def logout(response: Response) -> Response:
    response.delete_cookie(
        key=SESSION_COOKIE,
        path="/",
        httponly=True,
        secure=APP_ENV == "production",
        samesite="strict",
    )
    response.status_code = status.HTTP_204_NO_CONTENT
    return response


@app.post("/api/screenings", status_code=status.HTTP_201_CREATED, tags=["screenings"])
async def create_screening(
    screening_name: str = Form(..., min_length=2, max_length=160),
    resumes: list[UploadFile] = File(default=[]),
    job_description_text: str = Form(default=""),
    job_description_file: UploadFile | None = File(default=None),
    evaluation_criteria: str = Form(default="[]"),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if len(job_description_text) > MAX_JOB_DESCRIPTION_CHARS:
        raise HTTPException(status_code=422, detail=f"Keep the job description under {MAX_JOB_DESCRIPTION_CHARS:,} characters.")
    criteria = parse_evaluation_criteria(evaluation_criteria)

    clean_job_text = job_description_text.strip()
    if clean_job_text and count_words(clean_job_text) < MIN_JOB_DESCRIPTION_WORDS:
        raise HTTPException(status_code=422, detail=f"Job description must contain at least {MIN_JOB_DESCRIPTION_WORDS} words.")

    screening_name = screening_name.strip()
    if len(screening_name) < 2:
        raise HTTPException(status_code=422, detail="Enter a screening name.")

    uploaded_paths: list[Path] = []
    screening = Screening(
        id=str(uuid4()),
        name=screening_name,
        owner_id=user.id,
        job_description_text=clean_job_text,
        scoring_config=default_matching_weights().model_dump(),
    )
    screening.rubric = ScreeningRubric(criteria=criteria)
    try:
        if job_description_file is not None:
            filename, content = await read_upload(job_description_file, JOB_EXTENSIONS)
            extracted_text = extract_job_description(filename, content).strip()
            if len(extracted_text) > MAX_JOB_DESCRIPTION_CHARS:
                raise HTTPException(status_code=422, detail=f"Keep the job description under {MAX_JOB_DESCRIPTION_CHARS:,} characters.")
            if not clean_job_text:
                if count_words(extracted_text) < MIN_JOB_DESCRIPTION_WORDS:
                    raise HTTPException(
                        status_code=422,
                        detail=f"Job description file must contain at least {MIN_JOB_DESCRIPTION_WORDS} readable words. Try a searchable PDF or paste the text.",
                    )
                screening.job_description_text = extracted_text
            storage_name = f"{uuid4().hex}{Path(filename).suffix.lower()}"
            file_path = UPLOAD_DIR / storage_name
            file_path.write_bytes(content)
            uploaded_paths.append(file_path)
            screening.job_description_filename = filename
            screening.job_description_content_type = job_description_file.content_type
            screening.job_description_storage_name = storage_name

        await store_resume_uploads(screening, resumes, uploaded_paths)
        db.add(screening)
        db.commit()
        db.refresh(screening)
        return screening_payload(screening)
    except HTTPException:
        db.rollback()
        for file_path in uploaded_paths:
            file_path.unlink(missing_ok=True)
        raise
    except Exception:
        db.rollback()
        for file_path in uploaded_paths:
            file_path.unlink(missing_ok=True)
        raise


@app.get("/api/screenings", tags=["screenings"])
def list_screenings(
    skip: int = Query(default=0, ge=0),
    limit: int = Query(default=25, ge=1, le=100),
    search: str = Query(default="", max_length=120),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    conditions = [Screening.owner_id == user.id]
    clean_search = search.strip()
    if clean_search:
        pattern = f"%{clean_search}%"
        conditions.append(Screening.name.ilike(pattern))

    total = db.scalar(select(func.count()).select_from(Screening).where(*conditions)) or 0
    resume_stats = (
        select(
            Resume.screening_id.label("screening_id"),
            func.count(Resume.id).label("resume_count"),
            func.coalesce(
                func.sum(case((ResumeEvaluation.status == "completed", 1), else_=0)), 0
            ).label("completed_count"),
            func.coalesce(
                func.sum(
                    case(
                        (ResumeEvaluation.status.in_(("queued", "processing")), 1),
                        else_=0,
                    )
                ),
                0,
            ).label("active_count"),
            func.coalesce(
                func.sum(case((ResumeEvaluation.status == "failed", 1), else_=0)), 0
            ).label("failed_count"),
            func.coalesce(
                func.sum(case((ResumeEvaluation.status == "stopped", 1), else_=0)), 0
            ).label("stopped_count"),
            func.avg(func.coalesce(ResumeEvaluation.reviewer_override_score, ResumeEvaluation.score)).label("average_score"),
        )
        .outerjoin(ResumeEvaluation, ResumeEvaluation.resume_id == Resume.id)
        .group_by(Resume.screening_id)
        .subquery()
    )
    rows = db.execute(
        select(
            Screening,
            func.coalesce(resume_stats.c.resume_count, 0),
            func.coalesce(resume_stats.c.completed_count, 0),
            func.coalesce(resume_stats.c.active_count, 0),
            func.coalesce(resume_stats.c.failed_count, 0),
            func.coalesce(resume_stats.c.stopped_count, 0),
            resume_stats.c.average_score,
        )
        .outerjoin(resume_stats, resume_stats.c.screening_id == Screening.id)
        .where(*conditions)
        .order_by(Screening.created_at.desc(), Screening.id.desc())
        .offset(skip)
        .limit(limit)
    ).all()
    return {
        "total": total,
        "skip": skip,
        "limit": limit,
        "screenings": [
            {
                "id": screening.id,
                "name": screening.name,
                "created_at": iso_utc(screening.created_at),
                "resume_count": int(resume_count),
                "completed_count": int(completed_count),
                "active_count": int(active_count),
                "failed_count": int(failed_count),
                "stopped_count": int(stopped_count),
                "average_score": round(float(average_score)) if average_score is not None else None,
            }
            for screening, resume_count, completed_count, active_count, failed_count, stopped_count, average_score in rows
        ],
    }


@app.post("/api/screenings/{screening_id}/retry", tags=["screenings"])
def retry_screening(
    screening_id: str,
    mode: Literal["all", "failed"],
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening.id).where(
            Screening.id == screening_id,
            Screening.owner_id == user.id,
        )
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")

    owned_resume_ids = select(Resume.id).where(Resume.screening_id == screening_id).scalar_subquery()
    conditions = [ResumeEvaluation.resume_id.in_(owned_resume_ids)]
    if mode == "failed":
        conditions.append(ResumeEvaluation.status.in_(("failed", "stopped")))
    queued = db.execute(
        update(ResumeEvaluation)
        .where(*conditions)
        .values(
            status="queued",
            score=None,
            job_match_score=None,
            preferred_score=None,
            bonus_points=None,
            required_checks_passed=None,
            required_checks=None,
            job_evidence=None,
            explanation=None,
            assessments=None,
            matching_details=None,
            reviewer_override_score=None,
            reviewer_override_note=None,
            reviewer_id=None,
            reviewer_updated_at=None,
            error_message=None,
            attempts=0,
            available_at=None,
            updated_at=datetime.now(timezone.utc),
        )
    )
    db.commit()
    return {"queued_count": queued.rowcount or 0, "mode": mode}


@app.post("/api/screenings/{screening_id}/stop", tags=["screenings"])
def stop_screening(
    screening_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening.id).where(
            Screening.id == screening_id,
            Screening.owner_id == user.id,
        )
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")

    owned_resume_ids = select(Resume.id).where(Resume.screening_id == screening_id).scalar_subquery()
    stopped = db.execute(
        update(ResumeEvaluation)
        .where(
            ResumeEvaluation.resume_id.in_(owned_resume_ids),
            ResumeEvaluation.status.in_(("queued", "processing")),
        )
        .values(
            status="stopped",
            error_message="Stopped by recruiter.",
            available_at=None,
            updated_at=datetime.now(timezone.utc),
        )
    )
    db.commit()
    return {"stopped_count": stopped.rowcount or 0}


@app.post("/api/screenings/suggest-criteria", tags=["screenings"])
async def suggest_screening_criteria(
    job_description_text: str = Form(default=""),
    job_description_file: UploadFile | None = File(default=None),
    user: User = Depends(get_current_user),
):
    if len(job_description_text) > MAX_JOB_DESCRIPTION_CHARS:
        raise HTTPException(
            status_code=422,
            detail=f"Keep the job description under {MAX_JOB_DESCRIPTION_CHARS:,} characters.",
        )
    clean_job_text = job_description_text.strip()
    if job_description_file is not None:
        filename, content = await read_upload(job_description_file, JOB_EXTENSIONS)
        extracted_text = extract_job_description(filename, content).strip()
        if len(extracted_text) > MAX_JOB_DESCRIPTION_CHARS:
            raise HTTPException(
                status_code=422,
                detail=f"Keep the job description under {MAX_JOB_DESCRIPTION_CHARS:,} characters.",
            )
        if not clean_job_text:
            clean_job_text = extracted_text
    if count_words(clean_job_text) < MIN_JOB_DESCRIPTION_WORDS:
        raise HTTPException(
            status_code=422,
            detail=f"Job description must contain at least {MIN_JOB_DESCRIPTION_WORDS} words.",
        )
    try:
        return {"criteria": suggest_criteria(clean_job_text)}
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error


@app.post("/api/screenings/{screening_id}/resumes", tags=["resumes"])
async def upload_screening_resumes(
    screening_id: str,
    resumes: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening)
        .options(
            selectinload(Screening.resumes).selectinload(Resume.evaluation),
                selectinload(Screening.rubric),
        )
        .where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")
    uploaded_paths: list[Path] = []
    try:
        await store_resume_uploads(screening, resumes, uploaded_paths)
        db.commit()
        return screening_payload(screening)
    except HTTPException:
        db.rollback()
        for file_path in uploaded_paths:
            file_path.unlink(missing_ok=True)
        raise
    except Exception:
        db.rollback()
        for file_path in uploaded_paths:
            file_path.unlink(missing_ok=True)
        raise


@app.get("/api/screenings/{screening_id}", tags=["screenings"])
def get_screening(
    screening_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening)
        .options(
            selectinload(Screening.resumes).selectinload(Resume.evaluation),
                selectinload(Screening.rubric),
        )
        .where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")
    return screening_payload(screening)


@app.put("/api/screenings/{screening_id}/requirements", tags=["screenings"])
def update_screening_requirements(
    screening_id: str,
    payload: UpdateRequirementsRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening)
        .options(selectinload(Screening.resumes).selectinload(Resume.evaluation))
        .where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")
    requirements = [item.model_dump() for item in payload.requirements]
    requirement_ids = [item["id"] for item in requirements]
    if len(requirement_ids) != len(set(requirement_ids)):
        raise HTTPException(status_code=422, detail="Requirement IDs must be unique.")
    screening.job_requirements = requirements
    now = datetime.now(timezone.utc)
    for resume in screening.resumes:
        evaluation = resume.evaluation
        if evaluation is None:
            resume.evaluation = ResumeEvaluation(status="queued", updated_at=now)
            continue
        evaluation.status = "queued"
        evaluation.score = None
        evaluation.job_match_score = None
        evaluation.preferred_score = None
        evaluation.bonus_points = None
        evaluation.required_checks_passed = None
        evaluation.required_checks = None
        evaluation.job_evidence = None
        evaluation.explanation = None
        evaluation.assessments = None
        evaluation.matching_details = None
        evaluation.reviewer_override_score = None
        evaluation.reviewer_override_note = None
        evaluation.reviewer_id = None
        evaluation.reviewer_updated_at = None
        evaluation.error_message = None
        evaluation.attempts = 0
        evaluation.available_at = None
        evaluation.updated_at = now
    db.commit()
    return {"requirements": screening.job_requirements, "queued_count": len(screening.resumes)}


@app.patch("/api/screenings/{screening_id}", tags=["screenings"])
def update_screening(
    screening_id: str,
    payload: UpdateScreeningRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening)
        .options(
            selectinload(Screening.resumes).selectinload(Resume.evaluation),
            selectinload(Screening.rubric),
        )
        .where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")

    screening_name = payload.name.strip()
    if len(screening_name) < 2:
        raise HTTPException(status_code=422, detail="Enter a screening name.")
    job_description = payload.job_description_text.strip()
    if count_words(job_description) < MIN_JOB_DESCRIPTION_WORDS:
        raise HTTPException(
            status_code=422,
            detail=f"Job description must contain at least {MIN_JOB_DESCRIPTION_WORDS} words.",
        )
    criteria = (
        parse_evaluation_criteria(json.dumps(payload.criteria))
        if payload.criteria is not None
        else screening.rubric.criteria if screening.rubric else []
    )
    previous_job_file = (
        UPLOAD_DIR / screening.job_description_storage_name
        if screening.job_description_storage_name
        else None
    )

    screening.name = screening_name
    screening.job_description_text = job_description
    screening.job_requirements = None
    screening.job_description_filename = None
    screening.job_description_content_type = None
    screening.job_description_storage_name = None
    if screening.rubric is None:
        screening.rubric = ScreeningRubric(criteria=criteria)
    else:
        screening.rubric.criteria = criteria

    now = datetime.now(timezone.utc)
    for resume in screening.resumes:
        if resume.evaluation is None:
            resume.evaluation = ResumeEvaluation(status="queued", updated_at=now)
            continue
        resume.evaluation.status = "queued"
        resume.evaluation.score = None
        resume.evaluation.job_match_score = None
        resume.evaluation.preferred_score = None
        resume.evaluation.bonus_points = None
        resume.evaluation.required_checks_passed = None
        resume.evaluation.required_checks = None
        resume.evaluation.job_evidence = None
        resume.evaluation.explanation = None
        resume.evaluation.assessments = None
        resume.evaluation.matching_details = None
        resume.evaluation.error_message = None
        resume.evaluation.attempts = 0
        resume.evaluation.available_at = None
        resume.evaluation.updated_at = now

    db.commit()
    if previous_job_file is not None:
        previous_job_file.unlink(missing_ok=True)
    return screening_payload(screening)


@app.delete("/api/screenings/{screening_id}", tags=["screenings"])
def delete_screening(
    screening_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening)
        .options(selectinload(Screening.resumes))
        .where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None:
        raise HTTPException(status_code=404, detail="Screening not found.")

    stored_files = [UPLOAD_DIR / resume.storage_name for resume in screening.resumes]
    if screening.job_description_storage_name:
        stored_files.append(UPLOAD_DIR / screening.job_description_storage_name)
    db.delete(screening)
    db.commit()
    for file_path in stored_files:
        file_path.unlink(missing_ok=True)
    return {"deleted": True, "screening_id": screening_id}


@app.get("/api/resumes/{resume_id}/download", tags=["resumes"])
def download_resume(
    resume_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    resume = db.scalar(
        select(Resume)
        .join(Screening)
        .where(Resume.id == resume_id, Screening.owner_id == user.id)
    )
    if resume is None:
        raise HTTPException(status_code=404, detail="Resume not found.")
    file_path = UPLOAD_DIR / resume.storage_name
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="Stored resume file is missing.")
    return FileResponse(file_path, media_type=resume.content_type or "application/octet-stream", filename=resume.filename)


@app.patch("/api/resumes/{resume_id}/review", tags=["resumes"])
def review_resume_score(
    resume_id: str,
    payload: ReviewOverrideRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    evaluation = db.scalar(
        select(ResumeEvaluation)
        .join(Resume)
        .join(Screening)
        .where(Resume.id == resume_id, Screening.owner_id == user.id)
    )
    if evaluation is None:
        raise HTTPException(status_code=404, detail="Candidate evaluation not found.")
    if evaluation.status != "completed":
        raise HTTPException(status_code=409, detail="Wait for evaluation to complete before reviewing its score.")
    evaluation.reviewer_override_score = payload.score
    evaluation.reviewer_override_note = payload.note.strip()
    evaluation.reviewer_id = user.id
    evaluation.reviewer_updated_at = datetime.now(timezone.utc)
    db.commit()
    return {
        "resume_id": resume_id,
        "model_score": evaluation.score,
        "score": evaluation.reviewer_override_score,
        "reviewer_note": evaluation.reviewer_override_note,
        "reviewer_updated_at": iso_utc(evaluation.reviewer_updated_at),
    }


@app.get("/api/screenings/{screening_id}/job-description", tags=["screenings"])
def download_job_description(
    screening_id: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    screening = db.scalar(
        select(Screening).where(Screening.id == screening_id, Screening.owner_id == user.id)
    )
    if screening is None or not screening.job_description_storage_name:
        raise HTTPException(status_code=404, detail="Uploaded job description not found.")
    file_path = UPLOAD_DIR / screening.job_description_storage_name
    if not file_path.is_file():
        raise HTTPException(status_code=404, detail="Stored job description file is missing.")
    return FileResponse(
        file_path,
        media_type=screening.job_description_content_type or "application/octet-stream",
        filename=screening.job_description_filename,
    )


@app.post("/api/resumes/bulk-delete", tags=["resumes"])
def bulk_delete_resumes(
    payload: BulkDeleteResumes,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    resume_ids = list(dict.fromkeys(payload.resume_ids))
    if not resume_ids:
        raise HTTPException(status_code=422, detail="Select at least one resume to remove.")
    records = list(
        db.scalars(
            select(Resume)
            .join(Screening)
            .where(Resume.id.in_(resume_ids), Screening.owner_id == user.id)
        ).all()
    )
    for record in records:
        db.delete(record)
    db.commit()
    for record in records:
        (UPLOAD_DIR / record.storage_name).unlink(missing_ok=True)
    return {"deleted_count": len(records), "deleted_ids": [record.id for record in records]}
